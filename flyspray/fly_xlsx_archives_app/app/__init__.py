"""
Export data from flyspray to xlsx dashboards using SQLAlchemy ORM.
"""

import logging
from dataclasses import dataclass, field
from datetime import datetime, timezone
from os.path import join
from pathlib import Path
from typing import List, Optional
from zoneinfo import ZoneInfo

from app.config import DB_HOST, DB_PWD, DB_USER
from app.models import (
    FlysprayComment,
    FlysprayListCategory,
    FlysprayListStatus,
    FlysprayProject,
    FlysprayTask,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session, selectinload

logger = logging.getLogger(__name__)


# some dataclass
@dataclass
class TaskInfo:
    """
    data container for aggregated flyspray task details and associated comments.
    """
    task_id: int
    project: str
    summary: str
    desc: str
    status: str
    expl_team: str
    evt_nb: int
    open_dt: Optional[datetime] = None
    last_edit_dt: Optional[datetime] = None
    comments: List[FlysprayComment] = field(default_factory=list)


def parse_fly_timestamp(ts: Optional[int]) -> Optional[datetime]:
    """
    Convert a Unix UTC timestamp integer to a naive datetime localized to Europe/Paris.

    :param ts: Unix timestamp in seconds (or None/0).
    :return: A naive datetime object in Paris local time, or None if the input is empty.
    """
    if not ts:
        return None
    dt_utc = datetime.fromtimestamp(ts, tz=timezone.utc)
    return dt_utc.astimezone(ZoneInfo("Europe/Paris")).replace(tzinfo=None)


def db_request(database: str) -> List[TaskInfo]:
    """
    query flyspray tasks and associated comments from the specified mysql database using sqlalchemy.

    :param database: name of the mysql database to query.
    :return: list of TaskInfo instances containing extracted task details.
    """

    logger.info(f'get data from DB "{database}"')
    # create mysql engine (using pymysql driver under the hood)
    url = f'mysql+pymysql://{DB_USER}:{DB_PWD}@{DB_HOST}/{database}?charset=utf8mb4'
    engine = create_engine(url)

    try:
        with Session(engine) as session:
            # query tasks and eager-load comments relationship via selectinload
            stmt = (
                select(
                    FlysprayTask,
                    FlysprayProject.project_title,
                    FlysprayListCategory.category_name,
                    FlysprayListStatus.status_name,
                    func.count(FlysprayComment.comment_id).label('comments_nb'),
                )
                .options(
                    selectinload(FlysprayTask.comments).selectinload(FlysprayComment.user)
                )
                .outerjoin(
                    FlysprayProject,
                    FlysprayTask.project_id == FlysprayProject.project_id,
                )
                .outerjoin(
                    FlysprayListCategory,
                    FlysprayTask.product_category == FlysprayListCategory.category_id,
                )
                .outerjoin(
                    FlysprayListStatus,
                    FlysprayTask.item_status == FlysprayListStatus.status_id,
                )
                .outerjoin(
                    FlysprayComment,
                    FlysprayTask.task_id == FlysprayComment.task_id,
                )
                .group_by(FlysprayTask.task_id)
                .order_by(FlysprayTask.task_id.desc())
                .limit(100_000)
            )

            results = session.execute(stmt).all()

        open_task_l: List[TaskInfo] = []

        for row in results:
            task: FlysprayTask = row.FlysprayTask

            task_info = TaskInfo(
                task_id=task.task_id,
                project=row.project_title or '',
                summary=task.item_summary or '',
                desc=task.detailed_desc or '',
                status=row.status_name or '',
                expl_team=row.category_name or '',
                evt_nb=row.comments_nb,
                open_dt=parse_fly_timestamp(task.date_opened),
                last_edit_dt=parse_fly_timestamp(task.last_edited_time),
                comments=task.comments,
            )
            open_task_l.append(task_info)

        return open_task_l
    finally:
        engine.dispose()


def build_xlsx(file: str, task_l: list):
    """
    generate an excel spreadsheet dashboard from a list of taskinfo items.

    :param file: destination file path for the output xlsx workbook.
    :param task_l: list of taskinfo objects to write to the worksheet.
    """
    logger.info(f'build xlsx "{file}"')
    # Build XLSX workbook
    wb = Workbook()
    sheet = wb.active

    if sheet is None:
        raise RuntimeError('unable to find default xlsx sheet')

    sheet.title = "Feuil1"

    # Dimensions & Headings
    headers_and_widths = [
        ("Ticket", 10.0),
        ("Equipe DTS", 20.0),
        ("Exploitant", 30.0),
        ("Date et heure d'ouverture", 30.0),
        ("Date et heure modification", 30.0),
        ("Résumé", 75.0),
        ("Description", 75.0),
        ("Nb commentaires", 15.0),
        ("Commentaires", 80.0),
    ]

    for col_idx, (header, width) in enumerate(headers_and_widths, start=1):
        sheet.column_dimensions[get_column_letter(col_idx)].width = width
        sheet.cell(row=1, column=col_idx, value=header)

    # Append rows
    for row_idx, task_info in enumerate(task_l, start=2):
        # init column index
        col_idx = 0
        # column "Task ID"
        col_idx += 1
        sheet.cell(row=row_idx, column=col_idx, value=task_info.task_id)
        # column "Equipe DTS"
        col_idx += 1
        sheet.cell(row=row_idx, column=col_idx, value=task_info.project)
        # column "Exploitant"
        col_idx += 1
        sheet.cell(row=row_idx, column=col_idx, value=task_info.expl_team)
        # column "Date et Heure d'ouverture"
        col_idx += 1
        if task_info.open_dt:
            cell_open = sheet.cell(row=row_idx, column=col_idx, value=task_info.open_dt)
            cell_open.number_format = "dd/mm/yyyy hh:mm"
        # column "Date et Heure dernier événement"
        col_idx += 1
        if task_info.last_edit_dt:
            cell_evt = sheet.cell(row=row_idx, column=col_idx, value=task_info.last_edit_dt)
            cell_evt.number_format = "dd/mm/yyyy hh:mm"
        # column "Résumé"
        col_idx += 1
        sheet.cell(row=row_idx, column=col_idx, value=task_info.summary)
        # column "Description"
        col_idx += 1
        sheet.cell(row=row_idx, column=col_idx, value=task_info.desc)
        # column "Nb commentaires"
        col_idx += 1
        sheet.cell(row=row_idx, column=col_idx, value=str(task_info.evt_nb))
        # column "Commentaires"
        col_idx += 1
        cell_txt = ''
        for comment in task_info.comments:
            if comment.comment_text:
                comment_dt = parse_fly_timestamp(comment.last_edited_time)
                user_name = (
                    comment.user.real_name or comment.user.user_name
                    if comment.user
                    else 'unknown user'
                )
                head_mark = f'---- {comment_dt} créé par {user_name} ----'
                cell_txt += f'{head_mark}\r\n{comment.comment_text.strip()}\r\n\r\n'
        sheet.cell(row=row_idx, column=col_idx, value=cell_txt)

    # Format styles
    for col_idx in sheet.columns:
        for cell in col_idx:
            cell.alignment = Alignment(horizontal="center")

    tab = Table(
        displayName="Table1",
        ref=f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}",
    )
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium14", showRowStripes=True
    )
    sheet.add_table(tab)

    wb.save(file)


def main(pub_path: Path):
    """
    main orchestration function to iterate through target databases, fetch task details, and export excel reports.

    :param pub_path: directory path where the generated xlsx files will be saved.
    """

    databases = [("flyspray-tca", "tca"),
                 ("flyspray-tne", "tne"),
                 ("flyspray-trm", "trm"),
                 ("flyspray-tvs", "tvs")]

    for db_name, fly_id in databases:
        xlsx_file = join(pub_path, f"fly_{fly_id}_tasks.xlsx")
        task_info_l = db_request(db_name)
        build_xlsx(xlsx_file, task_info_l)
